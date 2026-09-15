import argparse
from html.parser import HTMLParser
import importlib.util
from io import BytesIO
import json
import os
from pathlib import Path
import sys
from urllib.error import HTTPError
from urllib.parse import urljoin
from urllib.request import build_opener, ProxyHandler, Request


class PageAssets(HTMLParser):
    def __init__(self):
        super().__init__()
        self.paths = set()

    def handle_starttag(self, tag, attrs):
        attributes = dict(attrs)
        path = attributes.get("src") if tag == "script" else attributes.get("href") if tag == "link" else None
        if path and path.startswith("/assets/"):
            self.paths.add(path)


def response(url, status=200, headers=None):
    opener = build_opener(ProxyHandler({}))
    try:
        result = opener.open(Request(url, headers=headers or {}), timeout=15)
    except HTTPError as error:
        result = error
    with result:
        assert result.status == status, f"{url}: expected {status}, got {result.status}"
        return result.read(2_000_000), result.headers


def verify(api, web):
    if hasattr(os, "geteuid"):
        assert os.geteuid() == 65532, "The API test image must run as its non-root user"
        assert sys.version_info[:3] == (3, 14, 7)
        assert not any(Path(path).exists() for path in ("/bin/sh", "/usr/bin/apt", "/sbin/apk", "/usr/bin/perl", "/usr/bin/mount"))
        assert all(importlib.util.find_spec(package) is None for package in ("pip", "setuptools", "wheel"))
    from cryptography.hazmat.primitives.asymmetric import rsa
    import jwt
    from openpyxl import Workbook, load_workbook
    from reportlab.pdfgen.canvas import Canvas

    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    signed = jwt.encode({"sub": "runtime-smoke"}, key, algorithm="RS256")
    assert jwt.decode(signed, key.public_key(), algorithms=["RS256"])["sub"] == "runtime-smoke"
    workbook = Workbook()
    workbook.active["A1"] = "runtime-smoke"
    workbook_data = BytesIO()
    workbook.save(workbook_data)
    workbook_data.seek(0)
    assert load_workbook(workbook_data).active["A1"].value == "runtime-smoke"
    pdf_data = BytesIO()
    canvas = Canvas(pdf_data)
    canvas.drawString(20, 20, "runtime-smoke")
    canvas.save()
    assert pdf_data.getvalue().startswith(b"%PDF-")
    payload, _ = response(urljoin(api, "/api/health"))
    assert json.loads(payload)["status"] == "ok"
    response(urljoin(api, "/api/auth/config"))
    response(urljoin(api, "/api/auth/me"), 401)
    response(urljoin(api, "/api/auth/me"), 401, {"X-MS-CLIENT-PRINCIPAL": "forged", "X-Meghkosha-User-Token": "forged"})
    payload, _ = response(urljoin(web, "/healthz"))
    assert payload.strip() == b"ok"
    payload, headers = response(web)
    assert b'id="root"' in payload
    assert headers["X-Content-Type-Options"] == "nosniff"
    assert headers["Referrer-Policy"] == "no-referrer"
    assert "frame-ancestors 'self'" in headers["Content-Security-Policy"]
    page = PageAssets()
    page.feed(payload.decode("utf-8"))
    assert page.paths, "The web image must contain compiled assets"
    for path in sorted(page.paths):
        asset, _ = response(urljoin(web, path))
        assert asset, f"Empty asset: {path}"
    response(urljoin(web, "/auth-callback.html"))
    response(urljoin(web, "/.env"), 403)
    response(urljoin(web, "/api/health"), 502)
    print(json.dumps({"result": "passed", "apiUser": 65532, "pythonVersion": sys.version.split()[0], "nativeDependencies": "passed", "compiledAssets": len(page.paths), "unsignedIdentity": "rejected", "unavailableHttpsUpstream": "rejected", "liveIdentityValidated": False}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--api", default="http://phase2-api:8000")
    parser.add_argument("--web", default="http://phase2-web:8080")
    arguments = parser.parse_args()
    verify(arguments.api, arguments.web)