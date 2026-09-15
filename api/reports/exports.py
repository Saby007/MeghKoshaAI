"""Deterministic stakeholder report artifacts built from persisted snapshots."""

from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from textwrap import shorten, wrap

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from reports.models import FinOpsActionState, ReportSnapshot

CUSTOM_REPORT_MODULES = {
    "summary",
    "subscriptions",
    "savings",
    "findings",
    "compute",
    "storage",
    "network",
    "sql",
    "ai",
    "governance",
    "actionPlan",
    "chargeback",
    "compliance",
}


@dataclass(frozen=True)
class ReportArtifact:
    content: bytes
    file_name: str
    media_type: str


def _money(value: float, currency: str) -> str:
    return f"{currency} {value:,.0f}"


def build_executive_pdf(snapshot: ReportSnapshot) -> ReportArtifact:
    report = snapshot.report
    metadata = report.report_metadata
    summary = report.executive_summary
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=letter, pageCompression=1)
    width, height = letter
    dark = HexColor("#171513")
    muted = HexColor("#5f5a57")
    orange = HexColor("#ee6018")
    green = HexColor("#4d7c43")
    line = HexColor("#d4cfca")

    pdf.setFillColor(dark)
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(42, height - 48, "MeghKoshaAI Executive Cost Assessment")
    pdf.setFillColor(orange)
    pdf.rect(42, height - 61, 62, 3, fill=1, stroke=0)
    pdf.setFillColor(muted)
    pdf.setFont("Helvetica", 8)
    pdf.drawRightString(width - 42, height - 45, f"Snapshot {snapshot.snapshot_id}")
    pdf.drawString(42, height - 78, f"Period: {metadata.period_start} to {metadata.period_end}")
    pdf.drawString(220, height - 78, f"Basis: {metadata.cost_basis}")
    pdf.drawRightString(width - 42, height - 78, f"Scope: {len(snapshot.subscription_ids)} subscription(s)")

    metrics = [
        ("Monthly spend", _money(summary.current_monthly_spend, metadata.currency), dark),
        ("Estimated wastage", _money(summary.estimated_wastage_month, metadata.currency), orange),
        ("Potential savings", _money(summary.potential_savings_month, metadata.currency), green),
        ("Other billed resources" if summary.idle_review_candidates is not None else "Active resources", f"{summary.active_resources:,}", dark),
        ("Confirmed idle" if summary.idle_review_candidates is not None else "Idle (legacy)", f"{summary.idle_resources:,}", orange),
        ("Advisor score", f"{report.advisor_score.score:.0f} / 100" if report.advisor_score.score is not None else "Unavailable", dark),
    ]
    top = height - 110
    card_width = (width - 84) / 3
    for index, (label, value, color) in enumerate(metrics):
        row, column = divmod(index, 3)
        x = 42 + column * card_width
        y = top - row * 61
        pdf.setStrokeColor(line)
        pdf.rect(x, y - 47, card_width - 8, 48, fill=0, stroke=1)
        pdf.setFillColor(muted)
        pdf.setFont("Helvetica", 7)
        pdf.drawString(x + 9, y - 13, label.upper())
        pdf.setFillColor(color)
        pdf.setFont("Helvetica-Bold", 15)
        pdf.drawString(x + 9, y - 34, value)

    y = top - 143
    pdf.setFillColor(dark)
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(42, y, "Where the money is going")
    y -= 16
    spend_total = max(summary.current_monthly_spend, 1)
    for item in report.spend_categories:
        if item.monthly_spend <= 0:
            continue
        pdf.setFillColor(muted)
        pdf.setFont("Helvetica", 8)
        pdf.drawString(42, y, item.category)
        pdf.drawRightString(205, y, _money(item.monthly_spend, metadata.currency))
        pdf.setFillColor(line)
        pdf.rect(215, y - 2, 130, 7, fill=1, stroke=0)
        pdf.setFillColor(orange)
        pdf.rect(215, y - 2, 130 * min(item.monthly_spend / spend_total, 1), 7, fill=1, stroke=0)
        y -= 15

    y -= 5
    pdf.setFillColor(dark)
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(42, y, "Top five savings and risk findings")
    y -= 17
    findings = report.prioritized_findings[:5]
    if not findings:
        pdf.setFillColor(muted)
        pdf.setFont("Helvetica", 8)
        pdf.drawString(42, y, "No prioritized savings or cost-at-risk findings were present.")
    for finding in findings:
        impact = finding.monthly_saving if finding.monthly_saving is not None else finding.monthly_cost_at_risk
        pdf.setFillColor(dark)
        pdf.setFont("Helvetica-Bold", 8)
        pdf.drawString(42, y, f"{finding.rank}. {shorten(finding.finding, width=72, placeholder='...')}")
        if impact is not None:
            pdf.drawRightString(width - 42, y, _money(impact, metadata.currency))
        y -= 11
        pdf.setFillColor(muted)
        pdf.setFont("Helvetica", 7)
        for text_line in wrap(finding.evidence, width=120)[:2]:
            pdf.drawString(54, y, text_line)
            y -= 9
        y -= 4

    pdf.setStrokeColor(line)
    pdf.line(42, 50, width - 42, 50)
    pdf.setFillColor(muted)
    pdf.setFont("Helvetica", 6.5)
    footer = (
        f"Source: {metadata.source} | Currency: {metadata.currency} | Generated: {snapshot.created_at} | "
        f"Completeness: {report.completeness.status}. Potential savings and cost at risk remain separate evidence classes."
    )
    for index, text_line in enumerate(wrap(footer, width=150)[:3]):
        pdf.drawString(42, 39 - index * 8, text_line)
    pdf.showPage()
    pdf.save()
    return ReportArtifact(
        content=output.getvalue(),
        file_name=f"MeghKoshaAI-Executive-{metadata.period or snapshot.snapshot_id}.pdf",
        media_type="application/pdf",
    )


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[object]]):
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(headers)
    for row in rows:
        sheet.append([ILLEGAL_CHARACTERS_RE.sub("", value)[:32767] if isinstance(value, str) else value for value in row])
    fill = PatternFill("solid", fgColor="24211F")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.font = Font(name="Arial", size=9, color="000000")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        sample = [len(str(sheet.cell(row=row, column=index).value or "")) for row in range(1, min(sheet.max_row, 100) + 1)]
        sheet.column_dimensions[get_column_letter(index)].width = min(max(max(sample, default=len(header)) + 2, 10), 48)
    return sheet


_SQL_ANALYSIS_HEADERS = [
    "Resource", "Resource ID", "Subscription ID", "Deployment model", "Rule", "Version", "Analysis", "Status",
    "Reason", "Required evidence", "Next steps", "After", "Savings estimate", "Configuration", "Workload evidence",
]


def _sql_analysis_rows(snapshot: ReportSnapshot, subscription_id: str | None = None) -> list[list[object]]:
    rows = []
    for category in snapshot.report.domains["sql"].categories:
        for line in category.lines:
            if subscription_id is not None and line.subscription_id != subscription_id:
                continue
            context = line.sql_context
            if context is None:
                continue
            for check in context.optimization_checks:
                rows.append([
                    line.resource_name, line.resource_id, line.subscription_id, context.deployment_model,
                    check.rule_id, check.rule_version, check.title, check.status, check.reason,
                    "\n".join(check.required_evidence), "\n".join(check.next_steps), ", ".join(check.depends_on),
                    "Unquantified", "\n".join(f"{key}: {value}" for key, value in context.configuration.items()),
                    context.workload_evidence.model_dump_json(by_alias=True) if context.workload_evidence else "Not collected",
                ])
    return rows


def build_full_xlsx(snapshot: ReportSnapshot, subscription_id: str | None = None) -> ReportArtifact:
    """Builds the full multi-tab workbook. When `subscription_id` is given, the Subscriptions,
    domain (Compute/Storage/Network/SQL/AI), AI Usage, and Governance sheets are filtered to
    that one subscription instead of the full combined scope; sheets that aggregate findings
    across subscriptions without a reliable per-subscription split are kept full and labeled
    accordingly rather than showing a fabricated single-subscription number.
    """
    report = snapshot.report
    metadata = report.report_metadata
    subscription_row = None
    if subscription_id is not None:
        subscription_row = next(
            (row for row in report.subscription_breakdown if row.subscription_id == subscription_id),
            None,
        )
        if subscription_row is None:
            raise ValueError(f"Subscription {subscription_id} is not part of this report scope")

    def _scoped_title(title: str) -> str:
        return title if subscription_row is None else f"{title} (All Subscriptions)"

    workbook = Workbook()
    workbook.remove(workbook.active)
    summary_rows = [
        ["Snapshot ID", snapshot.snapshot_id],
        ["Generated", snapshot.created_at],
        [
            "Scope",
            f"{subscription_row.subscription_name} ({subscription_row.subscription_id})"
            if subscription_row
            else f"{len(snapshot.subscription_ids)} subscription(s)",
        ],
        ["Period", metadata.period],
        ["Period start", metadata.period_start],
        ["Period end", metadata.period_end],
        ["Cost basis", metadata.cost_basis],
        ["Currency", metadata.currency],
        ["Source", metadata.source],
        ["Completeness", report.completeness.status],
        ["Monthly spend", subscription_row.current_spend if subscription_row else report.executive_summary.current_monthly_spend],
        ["Estimated wastage", subscription_row.total_waste if subscription_row else report.executive_summary.estimated_wastage_month],
        ["Potential savings / month", subscription_row.total_waste if subscription_row else report.executive_summary.potential_savings_month],
        ["Potential savings / year", (subscription_row.total_waste * 12) if subscription_row else report.executive_summary.potential_savings_year],
        ["Other billed resources (all subscriptions)" if report.executive_summary.idle_review_candidates is not None else "Active resources (legacy)", report.executive_summary.active_resources],
        ["Confirmed idle resources" if report.executive_summary.idle_review_candidates is not None else "Idle resources (legacy)", report.executive_summary.idle_resources],
        ["Idle review candidates", report.executive_summary.idle_review_candidates],
        ["Advisor score (all subscriptions)", report.advisor_score.score],
    ]
    summary = _sheet(workbook, "Summary", ["Measure", "Value"], summary_rows)
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 70

    _sheet(workbook, "Subscriptions", ["Subscription", "Subscription ID", f"Spend ({metadata.currency})", f"Verified saving ({metadata.currency})", "% recoverable"], [
        [row.subscription_name, row.subscription_id, row.current_spend, row.total_waste, row.pct_saved]
        for row in report.subscription_breakdown
        if subscription_id is None or row.subscription_id == subscription_id
    ])
    _sheet(workbook, _scoped_title("Savings Roadmap"), ["Opportunity", "Domain", "Impact type", f"Monthly ({metadata.currency})", f"Annual ({metadata.currency})", "Resources", "Immediate", "Risk", "Effort", "Action", "Prerequisites"], [
        [item.opportunity, item.domain, item.impact_type, item.monthly, item.annual, item.resources, item.immediate, item.risk, item.effort, item.recommended_action, " | ".join(item.prerequisites)]
        for item in report.savings_roadmap
    ])
    _sheet(workbook, _scoped_title("Prioritized Findings"), ["Rank", "Finding", "Category", "Impact type", f"Monthly saving ({metadata.currency})", f"Monthly cost at risk ({metadata.currency})", "Severity", "Evidence"], [
        [item.rank, item.finding, item.category, item.impact_type, item.monthly_saving, item.monthly_cost_at_risk, item.severity, item.evidence]
        for item in report.prioritized_findings
    ])

    for key, title in (("compute", "Compute"), ("storage", "Storage"), ("network", "Network"), ("sql", "Azure SQL"), ("ai", "AI Inventory")):
        rows = []
        for category in report.domains[key].categories:
            for line_item in category.lines:
                if subscription_id is not None and line_item.subscription_id != subscription_id:
                    continue
                rows.append([
                    category.display_name, category.impact_type, line_item.resource_name, line_item.resource_id,
                    line_item.subscription_name, line_item.subscription_id, line_item.monthly_cost,
                    line_item.evidence_type, line_item.confidence, line_item.detail,
                ])
        _sheet(workbook, title, ["Finding", "Impact type", "Resource", "Resource ID", "Subscription", "Subscription ID", f"Monthly cost ({metadata.currency})", "Evidence type", "Confidence", "Detail"], rows)

    sql_rows = _sql_analysis_rows(snapshot, subscription_id)
    if sql_rows:
        _sheet(workbook, "SQL Analysis", _SQL_ANALYSIS_HEADERS, sql_rows)
    ai_usage = report.ai_usage
    _sheet(workbook, "AI Usage", ["Account", "Deployment", "Model", "Version", "SKU", "Location", "Input tokens/day", "Output tokens/day", "Total tokens/day", f"Estimated cost/day ({metadata.currency})", "Trend", "Trend %", "Evidence"], [
        [item.account_name, item.deployment_name, item.model_name, item.model_version, item.sku_name, item.location, item.input_tokens_per_day, item.output_tokens_per_day, item.total_tokens_per_day, item.estimated_cost_day, item.trend_label, item.trend_percentage, item.evidence_status]
        for item in ai_usage.deployments
        if subscription_id is None or item.subscription_id == subscription_id
    ])
    _sheet(workbook, _scoped_title("AI Opportunities"), ["Priority", "Deployment", "Category", "Recommendation", "Evidence"], [
        [item.priority, item.deployment_name, item.category, item.recommendation, item.evidence]
        for item in ai_usage.opportunities
    ])
    _sheet(workbook, "Governance", ["Subscription", "Subscription ID", "Untagged resources", "% of untagged estate"], [
        [item.subscription_name, item.subscription_id, item.untagged_resources, item.pct_of_estate]
        for item in report.governance
        if subscription_id is None or item.subscription_id == subscription_id
    ])
    _sheet(workbook, _scoped_title("Action Plan"), ["Action", f"Saving/month ({metadata.currency})", "Prerequisite", "Affected subscriptions"], [
        [item.action, item.saving_month, item.prerequisite, ", ".join(value.subscription_name for value in item.affected_subscriptions)]
        for item in report.action_plan
    ])
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00;[Red](#,##0.00);-'
    output = BytesIO()
    workbook.save(output)
    file_name = f"MeghKoshaAI-Full-Assessment-{metadata.period or snapshot.snapshot_id}.xlsx"
    if subscription_row is not None:
        safe_name = re.sub(r"[^A-Za-z0-9_-]+", "-", subscription_row.subscription_name).strip("-") or subscription_row.subscription_id
        file_name = f"MeghKoshaAI-Full-Assessment-{safe_name}-{metadata.period or snapshot.snapshot_id}.xlsx"
    return ReportArtifact(
        content=output.getvalue(),
        file_name=file_name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def build_chargeback_xlsx(snapshot: ReportSnapshot) -> ReportArtifact:
    report = snapshot.report
    metadata = report.report_metadata
    chargeback = report.chargeback
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = _sheet(workbook, "Summary", ["Measure", "Value"], [
        ["Snapshot ID", snapshot.snapshot_id],
        ["Generated", snapshot.created_at],
        ["Period", metadata.period],
        ["Cost basis", metadata.cost_basis],
        ["Currency", metadata.currency],
        ["Source", metadata.source],
        ["Total monthly spend", report.executive_summary.current_monthly_spend],
        ["Team allocated cost", chargeback.allocated_cost],
        ["Team unallocated cost", chargeback.unallocated_cost],
        ["Team allocation coverage", chargeback.allocation_percentage],
        ["Status", chargeback.status],
    ])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 78
    allocation = _sheet(
        workbook,
        "Allocations",
        ["Dimension", "Tag value", f"Monthly cost ({metadata.currency})", "% of total", "Allocation status"],
        [
            [row.dimension, row.value, row.monthly_cost, row.pct_of_total, "Unallocated" if row.value == "Unallocated" else "Allocated"]
            for row in chargeback.rows
        ],
    )
    allocation.column_dimensions["A"].width = 18
    allocation.column_dimensions["B"].width = 32
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00;[Red](#,##0.00);-'
    for cell in summary[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
    output = BytesIO()
    workbook.save(output)
    return ReportArtifact(
        content=output.getvalue(),
        file_name=f"MeghKoshaAI-Chargeback-{metadata.period or snapshot.snapshot_id}.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def build_compliance_xlsx(snapshot: ReportSnapshot) -> ReportArtifact:
    report = snapshot.report
    metadata = report.report_metadata
    compliance = report.compliance
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "Summary", ["Measure", "Value"], [
        ["Snapshot ID", snapshot.snapshot_id],
        ["Generated", snapshot.created_at],
        ["Period", metadata.period],
        ["Scope", f"{len(snapshot.subscription_ids)} subscription(s)"],
        ["Status", compliance.status],
        ["Tagging denominator", "Azure Resource Graph resource count"],
        ["Policy denominator", "Current Azure Policy evaluation records; multiple evaluations can apply to one resource"],
    ])
    _sheet(workbook, "Tagging", ["Subscription", "Subscription ID", "Total resources", "Untagged resources", "% tagged"], [
        [row.subscription_name, row.subscription_id, row.total_resources, row.untagged_resources, row.tagging_percentage]
        for row in compliance.rows
    ])
    _sheet(workbook, "Policy Evaluations", [
        "Subscription", "Subscription ID", "Data available", "Policy assignments", "Compliant evaluations",
        "Noncompliant evaluations", "Conflict evaluations", "Exempt evaluations", "Not started evaluations",
        "Distinct noncompliant resources", "Evaluation compliance %", "Status",
    ], [
        [
            row.subscription_name, row.subscription_id, row.policy_data_available, row.policy_assignment_count,
            row.compliant_evaluations, row.non_compliant_evaluations, row.conflict_evaluations,
            row.exempt_evaluations, row.not_started_evaluations, row.non_compliant_resources,
            row.evaluation_compliance_percentage, row.status,
        ]
        for row in compliance.rows
    ])
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = '0.0%;[Red](0.0%);-'
    output = BytesIO()
    workbook.save(output)
    return ReportArtifact(
        content=output.getvalue(),
        file_name=f"MeghKoshaAI-Compliance-{metadata.period or snapshot.snapshot_id}.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _custom_title(period: str, label: str, multiple: bool) -> str:
    value = f"{period} {label}" if multiple else label
    return value[:31]


def build_custom_xlsx(snapshots: list[ReportSnapshot], modules: list[str]) -> ReportArtifact:
    if not snapshots:
        raise ValueError("At least one snapshot is required")
    selected = list(dict.fromkeys(modules))
    invalid = sorted(set(selected) - CUSTOM_REPORT_MODULES)
    if invalid or not selected:
        raise ValueError(f"Unsupported or empty custom report modules: {', '.join(invalid)}")
    scope = snapshots[0].subscription_ids
    currency = snapshots[0].report.report_metadata.currency
    cost_basis = snapshots[0].report.report_metadata.cost_basis
    if any(item.subscription_ids != scope for item in snapshots):
        raise ValueError("Custom report snapshots must use the same subscription scope")
    if any(item.report.report_metadata.currency != currency for item in snapshots):
        raise ValueError("Custom report snapshots must use one currency")
    if any(item.report.report_metadata.cost_basis != cost_basis for item in snapshots):
        raise ValueError("Custom report snapshots must use one cost basis")
    snapshots = sorted(snapshots, key=lambda item: (item.report.report_metadata.period_start, item.snapshot_id))
    workbook = Workbook()
    workbook.remove(workbook.active)
    multiple = len(snapshots) > 1
    domain_modules = {
        "compute": ("compute", "Compute"),
        "storage": ("storage", "Storage"),
        "network": ("network", "Network"),
        "sql": ("sql", "Azure SQL"),
        "ai": ("ai", "AI"),
    }
    for snapshot in snapshots:
        report = snapshot.report
        metadata = report.report_metadata
        period = metadata.period or snapshot.snapshot_id[:8]
        if "summary" in selected:
            _sheet(workbook, _custom_title(period, "Summary", multiple), ["Measure", "Value"], [
                ["Snapshot ID", snapshot.snapshot_id], ["Period", period], ["Period start", metadata.period_start],
                ["Period end", metadata.period_end], ["Subscriptions", len(snapshot.subscription_ids)],
                ["Currency", currency], ["Cost basis", cost_basis], ["Monthly spend", report.executive_summary.current_monthly_spend],
                ["Potential savings", report.executive_summary.potential_savings_month],
            ])
        if "subscriptions" in selected:
            _sheet(workbook, _custom_title(period, "Subscriptions", multiple), ["Subscription", "Subscription ID", f"Spend ({currency})", f"Verified saving ({currency})", "% recoverable"], [
                [row.subscription_name, row.subscription_id, row.current_spend, row.total_waste, row.pct_saved]
                for row in report.subscription_breakdown
            ])
        if "savings" in selected:
            _sheet(workbook, _custom_title(period, "Savings", multiple), ["Opportunity", "Domain", "Impact type", f"Monthly ({currency})", f"Annual ({currency})", "Resources", "Action"], [
                [item.opportunity, item.domain, item.impact_type, item.monthly, item.annual, item.resources, item.recommended_action]
                for item in report.savings_roadmap
            ])
        if "findings" in selected:
            _sheet(workbook, _custom_title(period, "Findings", multiple), ["Rank", "Finding", "Impact type", f"Monthly saving ({currency})", f"Cost at risk ({currency})", "Severity", "Evidence"], [
                [item.rank, item.finding, item.impact_type, item.monthly_saving, item.monthly_cost_at_risk, item.severity, item.evidence]
                for item in report.prioritized_findings
            ])
        for module, (domain_key, label) in domain_modules.items():
            if module not in selected:
                continue
            rows = []
            for category in report.domains[domain_key].categories:
                for line_item in category.lines:
                    rows.append([category.display_name, category.impact_type, line_item.resource_name, line_item.resource_id, line_item.subscription_name, line_item.monthly_cost, line_item.evidence_type, line_item.detail])
            _sheet(workbook, _custom_title(period, label, multiple), ["Finding", "Impact type", "Resource", "Resource ID", "Subscription", f"Monthly cost ({currency})", "Evidence", "Detail"], rows)
        if "sql" in selected:
            sql_rows = _sql_analysis_rows(snapshot)
            if sql_rows:
                _sheet(workbook, _custom_title(period, "SQL Analysis", multiple), _SQL_ANALYSIS_HEADERS, sql_rows)
        if "governance" in selected:
            _sheet(workbook, _custom_title(period, "Governance", multiple), ["Subscription", "Subscription ID", "Untagged resources", "% of estate"], [
                [item.subscription_name, item.subscription_id, item.untagged_resources, item.pct_of_estate] for item in report.governance
            ])
        if "actionPlan" in selected:
            _sheet(workbook, _custom_title(period, "Action Plan", multiple), ["Action", f"Saving/month ({currency})", "Prerequisite", "Subscriptions"], [
                [item.action, item.saving_month, item.prerequisite, ", ".join(value.subscription_name for value in item.affected_subscriptions)] for item in report.action_plan
            ])
        if "chargeback" in selected:
            _sheet(workbook, _custom_title(period, "Chargeback", multiple), ["Dimension", "Tag value", f"Monthly cost ({currency})", "% total", "Status"], [
                [row.dimension, row.value, row.monthly_cost, row.pct_of_total, "Unallocated" if row.value == "Unallocated" else "Allocated"] for row in report.chargeback.rows
            ])
        if "compliance" in selected:
            _sheet(workbook, _custom_title(period, "Compliance", multiple), ["Subscription", "Total resources", "Untagged", "% tagged", "Compliant evaluations", "Noncompliant evaluations", "Distinct noncompliant resources", "Assignments", "Evaluation compliance %", "Status"], [
                [row.subscription_name, row.total_resources, row.untagged_resources, row.tagging_percentage, row.compliant_evaluations, row.non_compliant_evaluations, row.non_compliant_resources, row.policy_assignment_count, row.evaluation_compliance_percentage, row.status]
                for row in report.compliance.rows
            ])
    output = BytesIO()
    workbook.save(output)
    first_period = snapshots[0].report.report_metadata.period or snapshots[0].snapshot_id[:8]
    last_period = snapshots[-1].report.report_metadata.period or snapshots[-1].snapshot_id[:8]
    period_label = first_period if first_period == last_period else f"{first_period}-to-{last_period}"
    return ReportArtifact(
        content=output.getvalue(),
        file_name=f"MeghKoshaAI-Custom-{period_label}.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def build_finops_monthly_xlsx(
    current: ReportSnapshot,
    previous: ReportSnapshot | None,
    states: list[FinOpsActionState],
) -> ReportArtifact:
    report = current.report
    metadata = report.report_metadata
    previous_report = previous.report if previous else None
    current_spend = report.executive_summary.current_monthly_spend
    previous_spend = previous_report.executive_summary.current_monthly_spend if previous_report else None
    spend_change = ((current_spend - previous_spend) / previous_spend) if previous_spend else None
    state_by_action = {state.action_id: state for state in states}
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "Summary", ["Measure", "Current", "Previous / context"], [
        ["Snapshot ID", current.snapshot_id, previous.snapshot_id if previous else "No prior cataloged period"],
        ["Period", metadata.period, previous_report.report_metadata.period if previous_report else None],
        ["Monthly spend", current_spend, previous_spend],
        ["Month-over-month spend change", spend_change, None],
        ["Potential savings / month", report.executive_summary.potential_savings_month, previous_report.executive_summary.potential_savings_month if previous_report else None],
        ["Tracked actions", len(states), None],
        ["Completed actions", sum(1 for state in states if state.status == "completed"), None],
        ["Realized saving / month", sum(state.realized_saving_month or 0 for state in states if state.status == "completed"), None],
        ["Evidence note", "Financial values use persisted complete report snapshots.", "Action states are audited workflow entries."],
    ])
    trend_rows = []
    for snapshot in ([previous] if previous else []) + [current]:
        if snapshot is None:
            continue
        value = snapshot.report
        trend_rows.append([
            value.report_metadata.period,
            value.report_metadata.period_start,
            value.report_metadata.period_end,
            value.executive_summary.current_monthly_spend,
            value.executive_summary.potential_savings_month,
            value.executive_summary.estimated_wastage_month,
            value.executive_summary.active_resources,
            value.executive_summary.idle_resources,
        ])
    _sheet(workbook, "Monthly Trend", ["Period", "Start", "End", f"Spend ({metadata.currency})", f"Potential savings ({metadata.currency})", f"Estimated wastage ({metadata.currency})", "Active resources", "Idle resources"], trend_rows)
    _sheet(workbook, "Actions", ["Action ID", "Action", "Subscriptions", f"Potential saving ({metadata.currency}/mo)", "Status", "Owner", "Due date", "Completed at", f"Realized saving ({metadata.currency}/mo)", "Note", "Updated at", "Updated by"], [
        [
            item.action_id, item.action, ", ".join(value.subscription_name for value in item.affected_subscriptions), item.saving_month,
            state_by_action[item.action_id].status if item.action_id in state_by_action else "open",
            state_by_action[item.action_id].owner if item.action_id in state_by_action else "",
            state_by_action[item.action_id].due_date if item.action_id in state_by_action else None,
            state_by_action[item.action_id].completed_at if item.action_id in state_by_action else None,
            state_by_action[item.action_id].realized_saving_month if item.action_id in state_by_action else None,
            state_by_action[item.action_id].note if item.action_id in state_by_action else "",
            state_by_action[item.action_id].updated_at if item.action_id in state_by_action else None,
            state_by_action[item.action_id].updated_by if item.action_id in state_by_action else None,
        ]
        for item in report.action_plan
    ])
    output = BytesIO()
    workbook.save(output)
    return ReportArtifact(
        content=output.getvalue(),
        file_name=f"MeghKoshaAI-FinOps-Monthly-{metadata.period or current.snapshot_id}.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )