from io import BytesIO
from datetime import date

from openpyxl import load_workbook
from pypdf import PdfReader

from brand import BRAND_NAME
from reports.builder import build_full_report
from reports.cost_details import build_cost_detail_export, build_cost_details
from anomalies.models import DailyCostRecord
from services.focus_history_reader import FocusHistoryData
from reports.exports import build_chargeback_xlsx, build_compliance_xlsx, build_custom_xlsx, build_executive_pdf, build_finops_monthly_xlsx, build_full_xlsx
from reports.models import ActionPlanItem, ChargebackSummary, ComplianceRow, ComplianceSummary, FinOpsActionState
from services.report_snapshots import build_snapshot


def _snapshot():
    report = build_full_report(
        subscription_ids=["sub-1"],
        subscription_names={"sub-1": "Sub One"},
        per_sub_spend={"sub-1": 100.0},
        resource_graph_rows={},
        cost_by_resource_id={},
        service_family_spend={"Compute": 100.0},
        service_spend={"Microsoft.Compute": 100.0},
        advisor_recommendations=[],
        untagged_counts={},
    )
    report = report.model_copy(update={
        "report_metadata": report.report_metadata.model_copy(update={
            "period": "2026-07",
            "period_start": "2026-07-01",
            "period_end": "2026-07-31",
            "currency": "USD",
        }),
    })
    return build_snapshot(report, ["sub-1"], 90, created_at="2026-08-17T08:00:00+00:00")


def test_cost_detail_workbook_filters_raw_tags_and_preserves_credits_and_literal_text():
    snapshot = _snapshot()
    history = FocusHistoryData("USD", "2026-07-01", "2026-07-02", 2, ["2026-07"], ["sub-1"], {"sub-1": "One"}, [
        DailyCostRecord("2026-07-01", "sub-1", "One", "Compute", "group", "/vm", "=formula", 24, tags={"app": "A", "owner": "Team"}),
        DailyCostRecord("2026-07-02", "sub-1", "One", "Compute", "group", "/vm", "=formula\x01", -2, tags={"app": "A", "owner": "Team"}),
        DailyCostRecord("2026-07-02", "sub-1", "One", "Storage", "group", "/disk", "disk", 100, tags={"app": "B"}),
    ])
    snapshot.report.cost_details = build_cost_details(history)
    artifact = build_cost_detail_export(snapshot, date(2026, 7, 2), date(2026, 7, 2), {"tagKey": "app", "tagValue": "A"})
    workbook = load_workbook(BytesIO(artifact.content), data_only=False)
    sheet = workbook["Daily resource costs"]
    assert sheet.max_row == 3
    assert all(sheet.cell(index, 4).value == "=formula" and sheet.cell(index, 4).data_type == "s" for index in (2, 3))
    assert sum(row[10] for row in sheet.iter_rows(min_row=2, values_only=True)) == -2
    assert sum(row[13] for row in sheet.iter_rows(min_row=2, values_only=True)) == 24
    assert sheet["P2"].value == "Not collected"
    assert workbook["Resources"]["C2"].value == -2
    assert workbook["Resources"]["D2"].value == 24
    assert "Tags" in workbook.sheetnames

    selected = build_cost_detail_export(snapshot, date(2026, 7, 1), date(2026, 7, 2), {"tagKey": "app", "tagValue": "A"}, selected_dates=[date(2026, 7, 2)])
    selected_workbook = load_workbook(BytesIO(selected.content))
    assert {row[0] for row in selected_workbook["Daily resource costs"].iter_rows(min_row=2, values_only=True)} == {"2026-07-02"}
    missing_owner = build_cost_detail_export(snapshot, date(2026, 7, 2), date(2026, 7, 2), {"requiredTagKeys": '["owner"]'})
    missing_workbook = load_workbook(BytesIO(missing_owner.content))
    assert missing_workbook["Resources"]["B2"].value == "/disk"
    assert missing_workbook["Resources"]["C2"].value == 100


def test_executive_pdf_is_one_page_with_snapshot_provenance():
    artifact = build_executive_pdf(_snapshot())

    assert artifact.content.startswith(b"%PDF")
    assert artifact.media_type == "application/pdf"
    reader = PdfReader(BytesIO(artifact.content))
    assert len(reader.pages) == 1
    text = reader.pages[0].extract_text()
    assert f"{BRAND_NAME} Executive Cost Assessment" in text
    assert "MONTHLY SPEND" in text.upper()
    assert "Snapshot" in text


def test_full_xlsx_has_expected_sheets_styles_and_no_formula_errors():
    artifact = build_full_xlsx(_snapshot())

    workbook = load_workbook(BytesIO(artifact.content), data_only=False)
    assert workbook.sheetnames == [
        "Summary", "Subscriptions", "Savings Roadmap", "Prioritized Findings", "Compute",
        "Storage", "Network", "Azure SQL", "AI Inventory", "AI Usage", "AI Opportunities",
        "Governance", "Action Plan",
    ]
    assert workbook["Summary"]["B2"].value == _snapshot().snapshot_id
    assert workbook["Subscriptions"].freeze_panes == "A2"
    assert workbook["Summary"]["A1"].font.name == "Arial"
    errors = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?"}
    assert not any(
        cell.value in errors
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )


def _multi_sub_snapshot():
    report = build_full_report(
        subscription_ids=["sub-1", "sub-2"],
        subscription_names={"sub-1": "Sub One", "sub-2": "Sub Two"},
        per_sub_spend={"sub-1": 1000.0, "sub-2": 500.0},
        resource_graph_rows={
            "unattached_disks": [
                {"id": "/sub/rg/disk1", "name": "disk1", "subscriptionId": "sub-1", "sizeGb": 128, "sku": "Premium_LRS"},
                {"id": "/sub/rg/disk2", "name": "disk2", "subscriptionId": "sub-2", "sizeGb": 64, "sku": "Standard_LRS"},
            ],
        },
        cost_by_resource_id={"/sub/rg/disk1": 10.0, "/sub/rg/disk2": 5.0},
        service_family_spend={"Storage": 15.0},
        advisor_recommendations=[],
        untagged_counts={
            "sub-1": {"total": 10, "untagged": 2},
            "sub-2": {"total": 5, "untagged": 1},
        },
    )
    report = report.model_copy(update={
        "report_metadata": report.report_metadata.model_copy(update={
            "period": "2026-07",
            "period_start": "2026-07-01",
            "period_end": "2026-07-31",
            "currency": "USD",
        }),
    })
    return build_snapshot(report, ["sub-1", "sub-2"], 90, created_at="2026-08-17T08:00:00+00:00")


def test_full_xlsx_scopes_to_one_subscription_and_labels_unscoped_sheets():
    snapshot = _multi_sub_snapshot()
    artifact = build_full_xlsx(snapshot, subscription_id="sub-1")

    workbook = load_workbook(BytesIO(artifact.content), data_only=False)
    assert "Savings Roadmap (All Subscriptions)" in workbook.sheetnames
    assert "Prioritized Findings (All Subscriptions)" in workbook.sheetnames
    assert "Action Plan (All Subscriptions)" in workbook.sheetnames
    assert "Savings Roadmap" not in workbook.sheetnames

    subscription_rows = list(workbook["Subscriptions"].iter_rows(min_row=2, values_only=True))
    assert [row[1] for row in subscription_rows] == ["sub-1"]

    storage_rows = list(workbook["Storage"].iter_rows(min_row=2, values_only=True))
    assert storage_rows and all(row[5] == "sub-1" for row in storage_rows)

    governance_rows = list(workbook["Governance"].iter_rows(min_row=2, values_only=True))
    assert [row[1] for row in governance_rows] == ["sub-1"]

    assert workbook["Summary"]["B4"].value == "Sub One (sub-1)"
    assert "sub-1" in artifact.file_name or "Sub-One" in artifact.file_name


def test_full_xlsx_rejects_a_subscription_outside_the_report_scope():
    import pytest

    with pytest.raises(ValueError):
        build_full_xlsx(_snapshot(), subscription_id="sub-not-in-scope")


def test_sql_analysis_workbooks_preserve_scope_evidence_and_literal_text():
    from findings.engine import build_report

    snapshot = _multi_sub_snapshot()
    rows = [{
        "id": f"/subscriptions/{subscription}/resourceGroups/rg/providers/Microsoft.Sql/servers/server/databases/db",
        "name": '=HYPERLINK("https://example.test","click")', "subscriptionId": subscription,
        "type": "microsoft.sql/servers/databases", "elasticPoolId": "", "licenseType": "LicenseIncluded",
    } for subscription in ("sub-1", "sub-2")]
    snapshot.report.domains["sql"].categories = build_report(["sub-1", "sub-2"], 0, {"sql_databases_and_pools": rows}, {}).tier_a_categories
    full = load_workbook(BytesIO(build_full_xlsx(snapshot, "sub-1").content), data_only=False)
    sheet = full["SQL Analysis"]
    assert sheet.max_row == 9
    assert {row[2] for row in sheet.iter_rows(min_row=2, values_only=True)} == {"sub-1"}
    assert sheet["A2"].value.startswith("=HYPERLINK")
    assert sheet["A2"].data_type == "s"
    assert sheet["M2"].value == "Unquantified"
    assert sheet["O2"].value == "Not collected"
    custom = load_workbook(BytesIO(build_custom_xlsx([snapshot], ["sql"]).content))
    assert custom["SQL Analysis"].max_row == 17
    assert set(row[4] for row in custom["SQL Analysis"].iter_rows(min_row=2, values_only=True)) == {f"SQL-R{index:02d}" for index in range(1, 9)}


def test_chargeback_xlsx_preserves_unallocated_rows_and_snapshot_id():
    snapshot = _snapshot()
    report = snapshot.report.model_copy(update={
        "chargeback": ChargebackSummary.model_validate({
            "available": True,
            "status": "Team allocations preserve missing tags as Unallocated.",
            "allocatedCost": 70,
            "unallocatedCost": 30,
            "allocationPercentage": 0.7,
            "rows": [
                {"dimension": "Team", "value": "Payments", "monthlyCost": 70, "pctOfTotal": 0.7},
                {"dimension": "Team", "value": "Unallocated", "monthlyCost": 30, "pctOfTotal": 0.3},
            ],
        }),
    })
    snapshot = snapshot.model_copy(update={"report": report})
    artifact = build_chargeback_xlsx(snapshot)

    workbook = load_workbook(BytesIO(artifact.content), data_only=False)
    assert workbook.sheetnames == ["Summary", "Allocations"]
    assert workbook["Summary"]["B2"].value == snapshot.snapshot_id
    rows = list(workbook["Allocations"].iter_rows(min_row=2, values_only=True))
    assert ("Team", "Unallocated", 30, 0.3, "Unallocated") in rows
    assert sum(row[2] for row in rows if row[0] == "Team") == 100


def test_compliance_xlsx_labels_resource_and_evaluation_denominators():
    snapshot = _snapshot()
    report = snapshot.report.model_copy(update={
        "compliance": ComplianceSummary(
            available=True,
            status="Policy evaluation data available.",
            rows=[ComplianceRow(
                subscriptionId="sub-1",
                subscriptionName="Sub One",
                totalResources=10,
                untaggedResources=2,
                taggingPercentage=0.8,
                compliantEvaluations=80,
                nonCompliantEvaluations=15,
                conflictEvaluations=5,
                exemptEvaluations=3,
                notStartedEvaluations=2,
                nonCompliantResources=4,
                policyAssignmentCount=7,
                evaluationCompliancePercentage=0.8,
                policyDataAvailable=True,
                status="Policy evaluation data available.",
            )],
        ),
    })
    artifact = build_compliance_xlsx(snapshot.model_copy(update={"report": report}))

    workbook = load_workbook(BytesIO(artifact.content), data_only=False)
    assert workbook.sheetnames == ["Summary", "Tagging", "Policy Evaluations"]
    assert "Resource Graph resource count" in workbook["Summary"]["B7"].value
    assert "evaluation records" in workbook["Summary"]["B8"].value
    assert workbook["Policy Evaluations"]["F2"].value == 15
    assert workbook["Policy Evaluations"]["J2"].value == 4


def test_custom_xlsx_uses_selected_modules_across_complete_periods():
    july = _snapshot()
    august_report = july.report.model_copy(update={
        "report_metadata": july.report.report_metadata.model_copy(update={
            "period": "2026-08",
            "period_start": "2026-08-01",
            "period_end": "2026-08-31",
        }),
    })
    august = july.model_copy(update={"snapshot_id": "august-snapshot", "report": august_report})
    artifact = build_custom_xlsx([august, july], ["summary", "subscriptions", "compliance"])

    workbook = load_workbook(BytesIO(artifact.content), data_only=False)
    assert workbook.sheetnames == [
        "2026-07 Summary", "2026-07 Subscriptions", "2026-07 Compliance",
        "2026-08 Summary", "2026-08 Subscriptions", "2026-08 Compliance",
    ]
    assert "2026-07-to-2026-08" in artifact.file_name


def test_custom_xlsx_rejects_mixed_subscription_scopes():
    first = _snapshot()
    second = first.model_copy(update={"subscription_ids": ["sub-2"]})

    import pytest

    with pytest.raises(ValueError, match="same subscription scope"):
        build_custom_xlsx([first, second], ["summary"])


def test_finops_monthly_xlsx_compares_periods_and_tracks_completed_action():
    current = _snapshot()
    action = ActionPlanItem(
        actionId="unattached_disks",
        action="Delete unattached disks",
        savingMonth=25,
        prerequisite="Owner validation",
        affectedSubscriptions=[{"subscriptionId": "sub-1", "subscriptionName": "Sub One"}],
    )
    current = current.model_copy(update={
        "report": current.report.model_copy(update={"action_plan": [action]}),
    })
    previous = current.model_copy(update={
        "snapshot_id": "previous-snapshot",
        "report": current.report.model_copy(update={
            "report_metadata": current.report.report_metadata.model_copy(update={"period": "2026-06"}),
            "executive_summary": current.report.executive_summary.model_copy(update={"current_monthly_spend": 80}),
        }),
    })
    state = FinOpsActionState(
        scopeHash=current.scope_hash,
        actionId="unattached_disks",
        status="completed",
        owner="owner@example.com",
        dueDate="2026-08-31",
        completedAt="2026-08-17T12:00:00Z",
        realizedSavingMonth=20,
        note="Validated",
        updatedAt="2026-08-17T12:00:00Z",
        updatedBy="tenant:user",
    )
    artifact = build_finops_monthly_xlsx(current, previous, [state])

    workbook = load_workbook(BytesIO(artifact.content), data_only=False)
    assert workbook.sheetnames == ["Summary", "Monthly Trend", "Actions"]
    assert workbook["Summary"]["B8"].value == 1
    assert workbook["Summary"]["B9"].value == 20
    assert workbook["Actions"]["E2"].value == "completed"
    assert workbook["Actions"]["I2"].value == 20